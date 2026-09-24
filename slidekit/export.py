"""Optional PDF/PPTX export. Content lives exclusively in decks/<date>/deck.py."""
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt
from pptx.oxml.xmlchemy import OxmlElement
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from xml.etree.ElementTree import Element, SubElement, tostring
from .fonts import font_family, font_path
WIDTH, HEIGHT = 1152, 648

def rgb(color):
    return tuple(int(color[i:i + 2], 16) / 255 for i in (0, 2, 4))

def export_deck(deck, output):
    font_paths = {bold: font_path(deck.meta, bold) for bold in (False, True)}
    family = font_family(deck.meta)
    designed_slides = deck.slides
    ROOT = output
    ROOT.mkdir(parents=True, exist_ok=True)
    # Keep SVGs as vectors in the website; render export stills with the same
    # browser engine so PDF/PPTX/Pillow can consume them without a new dependency.
    svg_paths = {deck.asset(picture['src']) for spec in designed_slides
                 for picture in spec.get('images', [])
                 if picture['src'].lower().endswith('.svg')}
    svg_previews = {}
    vector_previews = {}
    if svg_paths or any(s.get('vectors') for s in designed_slides):
        from playwright.sync_api import sync_playwright
        (ROOT / 'svg-previews').mkdir(exist_ok=True)
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page(device_scale_factor=2)
            for index, source in enumerate(sorted(svg_paths)):
                page.goto(source.as_uri())
                page.evaluate('document.fonts.ready')
                target = ROOT / 'svg-previews' / f'image-{index+1}.png'
                page.locator('svg').screenshot(path=str(target), omit_background=True)
                svg_previews[source] = target
            page.goto('about:blank')
            page.set_viewport_size({'width': WIDTH, 'height': HEIGHT})
            for spec in designed_slides:
                if not spec.get('vectors'):
                    continue
                svg = Element('svg', xmlns='http://www.w3.org/2000/svg',
                              width=str(WIDTH), height=str(HEIGHT), viewBox='0 0 1152 648')
                for tag, attributes in spec['vectors']:
                    SubElement(svg, tag, {k: str(v) for k, v in attributes.items()})
                page.set_content('<style>body{margin:0;background:transparent}</style>' + tostring(svg, encoding='unicode'))
                target = ROOT / 'svg-previews' / (spec['id'] + '.png')
                page.locator('svg').screenshot(path=str(target), omit_background=True)
                vector_previews[spec['id']] = target
            browser.close()
    STEM = ROOT / deck.meta["export_stem"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(WIDTH), Pt(HEIGHT)
    prs.core_properties.title = deck.meta['title'] + ' — ' + deck.meta['date_label']
    prs.core_properties.subject = 'Project meeting'
    pdf = canvas.Canvas(str(STEM.with_suffix('.pdf')), pagesize=(WIDTH, HEIGHT))
    pdf.setTitle(prs.core_properties.title)
    for bold, path in font_paths.items():
        pdfmetrics.registerFont(TTFont(family + ('-Bold' if bold else ''), str(path)))
    previews = []
    (ROOT / 'previews').mkdir(exist_ok=True)
    for index, spec in enumerate(designed_slides, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(spec['bg'])
        slide.notes_slide.notes_text_frame.text = spec['notes']
        preview = Image.new('RGB', (WIDTH * 2, HEIGHT * 2), '#' + spec['bg'])
        draw = ImageDraw.Draw(preview)
        pdf.setFillColorRGB(*rgb(spec['bg']))
        pdf.rect(0, 0, WIDTH, HEIGHT, fill=1, stroke=0)
        def draw_video():
            media = spec['video']
            movie, poster = deck.asset(media['src']), deck.asset(media['poster'])
            x, y, w, h = media.get('bounds', [0, 0, WIDTH, HEIGHT])
            slide.shapes.add_movie(str(movie), Pt(x), Pt(y), Pt(w), Pt(h),
                                   poster_frame_image=str(poster), mime_type='video/mp4')
            pdf.drawImage(str(poster), x, HEIGHT-y-h, width=w, height=h)
            still = Image.open(poster).convert('RGB').resize((w*2, h*2))
            dim = media.get('dim', 0)
            if dim:
                mask = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
                mask.fill.solid(); mask.fill.fore_color.rgb = RGBColor.from_string(spec['bg'])
                mask.line.fill.background()
                alpha = OxmlElement('a:alpha'); alpha.set('val', str(round(dim * 100000)))
                mask._element.spPr.solidFill.srgbClr.append(alpha)
                pdf.saveState(); pdf.setFillAlpha(dim); pdf.setFillColorRGB(*rgb(spec['bg']))
                pdf.rect(x, HEIGHT-y-h, w, h, fill=1, stroke=0); pdf.restoreState()
                still = Image.blend(still, Image.new('RGB', still.size, '#'+spec['bg']), dim)
            preview.paste(still, (x*2, y*2))

        if spec.get('video', {}).get('background'):
            draw_video()
        for x, y, width, height, color in spec.get('rects', []):
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(width), Pt(height))
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(color)
            shape.line.fill.background()
            pdf.setFillColorRGB(*rgb(color))
            pdf.rect(x, HEIGHT - y - height, width, height, fill=1, stroke=0)
            draw.rectangle((x * 2, y * 2, (x + width) * 2, (y + height) * 2), fill='#' + color)
        for picture_index, picture in enumerate(spec.get('images', [])):
            path = deck.asset(picture['src'])
            path = svg_previews.get(path, path)
            x, y, w, h = picture['bounds']
            still = Image.open(path).convert('RGBA')
            if picture.get('fit') == 'cover':
                position = picture.get('position', 'center').split()
                horizontal = {'left': 0, 'center': 0.5, 'right': 1}.get(position[0], None)
                if horizontal is None:
                    horizontal = float(position[0].rstrip('%')) / 100
                vertical_token = position[1] if len(position) > 1 else 'center'
                vertical = {'top': 0, 'center': 0.5, 'bottom': 1}.get(vertical_token, None)
                if vertical is None:
                    vertical = float(vertical_token.rstrip('%')) / 100
                still = ImageOps.fit(still, (round(w*2), round(h*2)), centering=(horizontal, vertical))
                zoom = picture.get('zoom', 1)
                if zoom > 1:
                    crop_w, crop_h = still.width / zoom, still.height / zoom
                    crop_x = (still.width - crop_w) * horizontal
                    crop_y = (still.height - crop_h) * vertical
                    still = still.crop((round(crop_x), round(crop_y),
                                        round(crop_x + crop_w), round(crop_y + crop_h)))
                    still = still.resize((round(w*2), round(h*2)), Image.Resampling.LANCZOS)
                crop_dir = ROOT / 'cropped-images'
                crop_dir.mkdir(exist_ok=True)
                path = crop_dir / f'{index}-{picture_index}.png'
                still.save(path)
            else:
                # Match the website's object-fit:contain; do not stretch scientific panels.
                ratio = min(w / still.width, h / still.height)
                fitted_w, fitted_h = still.width * ratio, still.height * ratio
                x, y = x + (w-fitted_w)/2, y + (h-fitted_h)/2
                w, h = fitted_w, fitted_h
                still = still.resize((round(w*2), round(h*2)))
            slide.shapes.add_picture(str(path), Pt(x), Pt(y), width=Pt(w), height=Pt(h))
            pdf.drawImage(str(path), x, HEIGHT-y-h, width=w, height=h, mask='auto')
            preview.paste(still, (round(x*2), round(y*2)), still)
        if spec['id'] in vector_previews:
            path = vector_previews[spec['id']]
            slide.shapes.add_picture(str(path), Pt(0), Pt(0), width=Pt(WIDTH), height=Pt(HEIGHT))
            pdf.drawImage(str(path), 0, 0, width=WIDTH, height=HEIGHT, mask='auto')
            still = Image.open(path).convert('RGBA').resize(preview.size)
            preview.paste(still, (0,0), still)
        for x, y, width, height, content, size, bold, color in spec['text']:
            font_name = family + ('-Bold' if bold else '')
            assert size >= 30
            assert pdfmetrics.stringWidth(content, font_name, size) <= width, (index, content)
            assert x + width <= WIDTH and y + height <= HEIGHT, (index, content)
            box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(width), Pt(height))
            frame = box.text_frame
            frame.word_wrap = False
            frame.margin_left = frame.margin_right = Pt(0)
            frame.margin_top = frame.margin_bottom = Pt(0)
            p = frame.paragraphs[0]
            p.space_before = p.space_after = Pt(0)
            p.font.name, p.font.size, p.font.bold = family, Pt(size), bold
            p.font.color.rgb = RGBColor.from_string(color)
            p.text = content
            ascent = pdfmetrics.getAscent(font_name) * size / 1000
            pdf.setFont(font_name, size)
            pdf.setFillColorRGB(*rgb(color))
            pdf.drawString(x, HEIGHT - y - ascent, content)
            font = ImageFont.truetype(str(font_paths[bold]), size * 2)
            draw.text((x * 2, (y + ascent) * 2), content, font=font, fill='#' + color, anchor='ls')
        if spec.get('widget') or spec.get('scene'):
            snapshot = ROOT / 'widget-previews' / (spec['id'] + '.jpg')
            if not snapshot.exists():
                raise FileNotFoundError('Capture widget previews before exporting: ' + str(snapshot))
            x,y,w,h = spec.get('scene', {}).get('bounds', [48,92,1056,548])
            slide.shapes.add_picture(str(snapshot), Pt(x), Pt(y), width=Pt(w), height=Pt(h))
            pdf.drawImage(str(snapshot), x, HEIGHT-y-h, width=w, height=h)
            preview.paste(Image.open(snapshot).resize((w*2,h*2)), (x*2,y*2))
        if spec.get('video') and not spec['video'].get('background'):
            draw_video()
        preview.save(ROOT / 'previews' / f'slide-{index:02d}.png')
        if index == 1:
            preview.save(ROOT / 'recap-preview.png')
        previews.append(preview)
        pdf.showPage()
    prs.save(STEM.with_suffix('.pptx'))
    pdf.save()
    sheet = Image.new('RGB', (1152, ((len(previews) + 1) // 2) * 348), '#DCE2E8')
    for index, preview in enumerate(previews):
        thumb = preview.resize((560, 315), Image.Resampling.LANCZOS)
        sheet.paste(thumb, (8 + (index % 2) * 576, 8 + (index // 2) * 348))
    sheet.save(ROOT / 'deck-overview.png')
    print(f'Created {len(designed_slides)} slides: editable PPTX, PDF, PNG previews. Text bounds verified.')
